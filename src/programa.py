"""
PREDIMIN - Evaluacion de un programa de voladuras (varios disparos a la vez).

Flujo en la aplicacion:
  1. Descargar la plantilla Excel (plantilla()).
  2. Llenar una fila por disparo programado.
  3. Subir el archivo: evaluar_programa() predice cada disparo, asigna el
     nivel de alerta y la medida principal, y ordena por riesgo.
  4. Descargar el resultado con colores por nivel (a_excel()).
"""

from io import BytesIO

import numpy as np
import pandas as pd

from recomendador import ORDEN, prediccion_red_neuronal, recomendar

# Columna de la plantilla -> clave interna del modelo
COLUMNAS = {
    "Código del disparo": "codigo",
    "Fecha": "fecha",
    "Hora (0-23)": "hora",
    "Número de taladros": "numero_taladros",
    "Tonelaje fracturado (t)": "tonelaje_tm",
    "ANFO (kg)": "anfo_kg",
    "Emulsión (kg)": "emulsion_kg",
    "Disparos en el día": "n_eventos",
    "Retardo entre taladros (ms)": "tiempo_taladro_ms",
    "Retardo entre filas (ms)": "tiempo_fila_ms",
    "Humedad relativa (%)": "humedad_relativa_pct",
    "Velocidad del viento (m/s)": "velocidad_viento_ms",
    "Dirección del viento (°)": "direccion_viento_grados",
    "Precipitación (mm)": "precipitacion_mm",
}
OBLIGATORIAS = [c for c in COLUMNAS if c != "Código del disparo"]

TEXTO_NIVEL = {"VERDE": "Bajo", "AMARILLO": "Moderado", "NARANJA": "Alto", "ROJO": "Crítico"}
COLOR_EXCEL = {"VERDE": "D5F0D5", "AMARILLO": "FDEDC4", "NARANJA": "F9D6C6", "ROJO": "F4C7C7"}


def plantilla() -> bytes:
    """Excel de ejemplo con cinco disparos (valores tipicos de los registros)."""
    ej = pd.DataFrame([
        ["LA-0101", "2026-10-05", 12, 130, 64729, 2760, 11277, 2, 17, 182, 35, 2.5, 306, 0.0],
        ["LA-0102", "2026-10-05", 13, 95, 48000, 2100, 8600, 2, 17, 182, 38, 2.1, 300, 0.0],
        ["LA-0103", "2026-10-06", 12, 160, 82000, 3400, 13900, 1, 25, 200, 62, 1.4, 290, 0.0],
        ["LA-0104", "2026-10-07", 7, 120, 60000, 2500, 10500, 1, 17, 182, 85, 0.3, 160, 1.2],
        ["LA-0105", "2026-10-08", 12, 210, 105000, 4300, 18200, 3, 42, 224, 28, 3.4, 315, 0.0],
    ], columns=list(COLUMNAS))
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        ej.to_excel(w, sheet_name="Programa", index=False)
        _formato(w.sheets["Programa"])
        pd.DataFrame({"Columna": list(COLUMNAS), "Descripción": [
            "Identificador libre (opcional)", "Fecha programada (AAAA-MM-DD)",
            "Hora del disparo, 0 a 23", "Taladros a disparar", "Tonelaje a fracturar",
            "Kilogramos de ANFO", "Kilogramos de emulsión", "Disparos programados ese día",
            "Retardo entre taladros", "Retardo entre filas",
            "Pronóstico o medición de humedad relativa", "Pronóstico de velocidad del viento",
            "Dirección del viento (0-360°, de donde viene)", "Precipitación esperada"]}
        ).to_excel(w, sheet_name="Instrucciones", index=False)
        _formato(w.sheets["Instrucciones"])
    return buf.getvalue()


def leer(archivo) -> pd.DataFrame:
    df = pd.read_excel(archivo, sheet_name=0)
    df.columns = [str(c).strip() for c in df.columns]
    faltan = [c for c in OBLIGATORIAS if c not in df.columns]
    if faltan:
        raise ValueError("Faltan columnas en el archivo: " + ", ".join(faltan)
                         + ". Use la plantilla de PREDIMIN.")
    return df.dropna(how="all")


def evaluar_programa(df: pd.DataFrame) -> pd.DataFrame:
    """Predice cada disparo y devuelve la tabla ordenada de mayor a menor riesgo."""
    filas = []
    for i, fila in df.reset_index(drop=True).iterrows():
        d = {COLUMNAS[c]: fila[c] for c in COLUMNAS if c in df.columns}
        fecha = pd.to_datetime(d.get("fecha"), errors="coerce")
        try:
            evento = {k: float(d[k]) for k in COLUMNAS.values() if k not in ("codigo", "fecha")}
            if pd.isna(fecha) or any(np.isnan(v) for v in evento.values()):
                raise ValueError("hay celdas vacías o no numéricas")
        except (ValueError, TypeError) as e:
            filas.append({"Código": d.get("codigo") or f"Fila {i + 2}", "Observación":
                          f"No evaluado: {e}"})
            continue
        evento["mes"] = fecha.month
        evento["explosivo_total_kg"] = evento["anfo_kg"] + evento["emulsion_kg"]
        res = recomendar(evento)
        p = res["prediccion"]["pm10_ugm3"]
        rn = prediccion_red_neuronal(evento)
        medidas = [r["nombre"] for r in res["recomendaciones"]]
        filas.append({
            "Código": d.get("codigo") if pd.notna(d.get("codigo")) else f"Fila {i + 2}",
            "Fecha": fecha.date(), "Hora": int(evento["hora"]),
            "PM10 esperado (µg/m³)": p["valor"],
            "Prob. > 45 µg/m³ (%)": (None if p["probabilidad_superar_oms"] is None
                                     else round(p["probabilidad_superar_oms"] * 100)),
            "Red neuronal PM10 (µg/m³)": rn["valor"] if rn else None,
            "Nivel": res["nivel_alerta"],
            "Riesgo": TEXTO_NIVEL[res["nivel_alerta"]],
            "Medida principal": medidas[0] if medidas else "",
            "Otras medidas": "; ".join(medidas[1:]),
            "Observación": "",
        })
    out = pd.DataFrame(filas)
    if "Nivel" in out:
        out["_orden"] = out["Nivel"].map(ORDEN).fillna(-1)
        out = (out.sort_values(["_orden", "PM10 esperado (µg/m³)"], ascending=False)
               .drop(columns="_orden").reset_index(drop=True))
    return out


def a_excel(res: pd.DataFrame) -> bytes:
    """Resultado en Excel, con la fila coloreada segun el nivel de alerta."""
    from openpyxl.styles import PatternFill

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        res.drop(columns=["Nivel"], errors="ignore").to_excel(w, sheet_name="Evaluación",
                                                               index=False)
        ws = w.sheets["Evaluación"]
        _formato(ws)
        for i, nivel in enumerate(res.get("Nivel", []), start=2):
            if nivel in COLOR_EXCEL:
                relleno = PatternFill("solid", start_color=COLOR_EXCEL[nivel])
                for celda in ws[i]:
                    celda.fill = relleno
    return buf.getvalue()


def _formato(ws):
    from openpyxl.styles import Alignment, Font, PatternFill

    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col in ws.columns:
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, ancho + 2), 45)
    ws.freeze_panes = "A2"


if __name__ == "__main__":
    r = evaluar_programa(leer(BytesIO(plantilla())))
    print(r.to_string(index=False))
