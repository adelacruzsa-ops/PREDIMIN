"""
PREDIMIN - Reporte de resultados para la tesis.

Reune en un solo Excel (outputs/RESULTADOS_TESIS.xlsx) todas las tablas que
genera el sistema, con formato listo para copiar al documento:

  1_Depuracion        registros descartados y motivo (metodologia)
  2_Descriptivos      estadisticos de cada variable (media, desv., min, max...)
  3_Correlaciones     Spearman con PM10 (objetivo especifico 1)
  4_Modelos           comparacion de los 7 modelos de regresion (objetivo 2)
  5_Clasificacion     modelos de riesgo de superar 45 ug/m3
  6_Importancia_SHAP  variables mas influyentes (objetivo 3)
  7_Escenarios        reduccion estimada para un evento tipico (objetivo 5)
  8_Hiperparametros   configuracion final de cada modelo
  9_Red_Neuronal      arquitectura y desempeño de la red neuronal (ensamble)
  10_RN_Importancia   importancia de variables para la red (permutacion)

Uso:
    python src/reporte_tesis.py
"""

import json

import pandas as pd

from config import DIR_SALIDAS, ETIQUETAS as NOMBRES, OBJETIVOS
from importar_excel import importar
from preprocesamiento import columnas_modelo, preparar



def _leer(nombre):
    ruta = DIR_SALIDAS / nombre
    return pd.read_csv(ruta) if ruta.exists() else None


def _formatear(ws):
    from openpyxl.styles import Alignment, Font, PatternFill
    relleno = PatternFill("solid", start_color="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col in ws.columns:
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, ancho + 2), 50)
    ws.freeze_panes = "A2"


def main():
    hojas = {}

    raw = importar(verbose=False)
    hojas["1_Depuracion"] = pd.DataFrame(
        list(raw.attrs["log_depuracion"].items()), columns=["Concepto", "Registros"])

    df = preparar(verbose=False)
    cols = columnas_modelo() + OBJETIVOS
    desc = df[cols].describe().T.round(3)
    desc.insert(0, "Variable", [NOMBRES.get(c, c) for c in desc.index])
    hojas["2_Descriptivos"] = desc.reset_index(drop=True)

    corr = _leer("correlaciones_pm10.csv")
    if corr is not None:
        corr.insert(0, "Nombre", corr["variable"].map(lambda c: NOMBRES.get(c, c)))
        hojas["3_Correlaciones"] = corr

    for obj in OBJETIVOS:
        t = _leer(f"comparacion_modelos_{obj}.csv")
        if t is not None:
            hojas["4_Modelos"] = t
        t = _leer(f"comparacion_clasificadores_{obj}.csv")
        if t is not None:
            hojas["5_Clasificacion"] = t
        t = _leer(f"shap_importancia_{obj}.csv")
        if t is not None:
            t.insert(0, "Nombre", t["variable"].map(lambda c: NOMBRES.get(c, c)))
            hojas["6_Importancia_SHAP"] = t

    try:
        from recomendador import EVENTO_EJEMPLO, evaluar_escenarios
        hojas["7_Escenarios"] = evaluar_escenarios(EVENTO_EJEMPLO)
    except Exception as e:
        print(f"[aviso] No se calcularon escenarios (entrene primero): {e}")

    for obj in OBJETIVOS:
        ruta = DIR_SALIDAS / f"hiperparametros_{obj}.json"
        if ruta.exists():
            hp = json.loads(ruta.read_text(encoding="utf-8"))
            hojas["8_Hiperparametros"] = pd.DataFrame(
                [{"Modelo": m, "Hiperparametros": json.dumps(p, ensure_ascii=False)}
                 for m, p in hp.items()])

    ruta = DIR_SALIDAS / "red_neuronal_resumen.json"
    if ruta.exists():
        rn = json.loads(ruta.read_text(encoding="utf-8"))
        filas = [{"Seccion": sec, "Indicador": k,
                  "Valor": json.dumps(v) if isinstance(v, list) else v}
                 for sec in ("arquitectura", "regresion", "clasificacion")
                 for k, v in rn[sec].items()]
        hojas["9_Red_Neuronal"] = pd.DataFrame(filas)
    t = _leer("red_neuronal_importancia_permutacion.csv")
    if t is not None:
        t.insert(0, "Nombre", t["variable"].map(lambda c: NOMBRES.get(c, c)))
        hojas["10_RN_Importancia"] = t

    salida = DIR_SALIDAS / "RESULTADOS_TESIS.xlsx"
    with pd.ExcelWriter(salida, engine="openpyxl") as w:
        for nombre, tabla in hojas.items():
            tabla.to_excel(w, sheet_name=nombre, index=False)
            _formatear(w.sheets[nombre])
    print(f"[OK] Reporte con {len(hojas)} hojas -> {salida}")


if __name__ == "__main__":
    main()
